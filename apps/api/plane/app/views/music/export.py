# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

"""Catalog export: a registry of columns the user picks from.

Every column has a key (what the web sends in ``columns=``), a group (how the
export panel lists it), an English header and a value function. Headers are
the ones the catalog importer maps on its own (IMPORT_ALIASES /
_infer_mapping), so an exported file can be imported back: "Track Name",
"Isrc", "Track Display Artist", "Performing Contributors", "Writers"…

Credits are written the way distributor files write them, one entry per
person and role: ``Sergio "Keko" Vazquez <MixingEngineer>, Sergio "Keko"
Vazquez <MasteringEngineer>``.

Without ``columns`` the export keeps the original 19 columns and headers, so
older links and habits keep working. Everything a column needs is fetched
once for all songs (no query per song).
"""

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Callable

from django.db import connection
from django.db.models import Prefetch, Q
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from plane.db.models import (
    ContractWorkLink,
    EntityTagLink,
    MusicCompany,
    MusicCredit,
    MusicDistribution,
    MusicLink,
    MusicRelease,
    MusicReleaseArtist,
    MusicReleaseTrack,
    MusicTrack,
    MusicTrackGenre,
    MusicTrackImport,
)
from plane.utils.music_credits import (
    PERFORMER_CATEGORY,
    CreditTagger,
    RoleCatalog,
    format_tagged,
    workspace_credit_tagger,
    workspace_role_catalog,
)

# Groups in the order the export panel shows them.
GROUPS = (
    "song",  # Canción
    "identifiers",  # Identificadores
    "release",  # Lanzamiento
    "credits",  # Artistas y créditos
    "companies",  # Compañías y porcentajes
    "media",  # Videos y enlaces
    "rights",  # Derechos y textos
    "contracts",  # Contratos
    "import",  # Importación y datos extra
    "tags",  # Etiquetas
)

# Built-in roles the "Writers" column holds (the rest of the writer-like
# roles a file brings, e.g. "Librettist", stay where the file had them).
WRITER_ROLES = (MusicCredit.Role.WRITER, MusicCredit.Role.AUTHOR, MusicCredit.Role.LYRICIST, MusicCredit.Role.COMPOSER)
ARTIST_ROLES = (MusicCredit.Role.PRIMARY_ARTIST, MusicCredit.Role.FEATURED_ARTIST)

# The key of the "Datos extra" column: one column per metadata.extra key.
EXTRA = "extra"
EXTRA_PREFIX = "Extra: "
# Excel refuses longer cells
EXCEL_CELL_LIMIT = 32767


@dataclass
class ExportContext:
    """What the value functions read besides the song itself."""

    tagger: CreditTagger
    catalog: RoleCatalog
    contracts: dict = field(default_factory=dict)  # track id -> [(contract, rights_type)]
    tags: dict = field(default_factory=dict)  # track id -> [tag names]
    extra_keys: list = field(default_factory=list)


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    group: str
    value: Callable[[ExportContext, MusicTrack], Any]
    # Relations to fetch once for every song (see _prefetch)
    needs: tuple = ()


# ── Value helpers ──────────────────────────────────────────────────────────


def _join(values, separator="; "):
    return separator.join(str(value) for value in values if value not in (None, ""))


def _unique(values):
    seen, result = set(), []
    for value in values:
        if value in (None, ""):
            continue
        marker = str(value).lower() if isinstance(value, str) else value
        if marker not in seen:
            seen.add(marker)
            result.append(value)
    return result


def _one_or_many(values):
    """A single value keeps its type (a date stays a date in Excel); several
    distinct ones are joined."""
    values = _unique(values)
    if not values:
        return ""
    if len(values) == 1:
        return values[0]
    return _join(_text(value) for value in values)


def _aligned(values):
    """Parallel lists (videos, contracts): one slot per item, "-" for blanks,
    so the n-th value of each column belongs to the same item."""
    values = list(values)
    if not values:
        return ""
    if len(values) == 1:
        return values[0] if values[0] not in (None, "") else ""
    return "; ".join(_text(value) if value not in (None, "") else "-" for value in values)


def _text(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value)


def _number(value):
    """50.0000 -> "50", 12.5000 -> "12.5"."""
    if value is None:
        return ""
    if isinstance(value, Decimal):
        text = format(value.normalize(), "f")
        return text if text != "-0" else "0"
    return value


def _duration(ms):
    """3:16, or 1:02:05 past an hour (what the importer reads back)."""
    if ms in (None, ""):
        return ""
    seconds = round(ms / 1000)
    hours, rest = divmod(seconds, 3600)
    minutes, secs = divmod(rest, 60)
    return f"{hours}:{minutes:02d}:{secs:02d}" if hours else f"{minutes}:{secs:02d}"


def _clock(ms):
    """00:00:47, like distributor files write the preview start."""
    if ms in (None, ""):
        return ""
    seconds = round(ms / 1000)
    hours, rest = divmod(seconds, 3600)
    minutes, secs = divmod(rest, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _yes_no(value):
    return "Yes" if value else "No"


def _label(choices, value):
    return dict(choices).get(value, value) if value else ""


def _local_date(value):
    return value.date() if isinstance(value, datetime) else value


# ── Relations ──────────────────────────────────────────────────────────────


def _credits(track):
    return list(track.credits.all())


def _tagged(ctx, credits):
    return format_tagged(
        (credit.party.display_name, tag)
        for credit in credits
        for tag in ctx.tagger.tags_for(credit.role, credit.role_detail)
    )


def _names(credits, roles):
    return ", ".join(_unique(credit.party.display_name for credit in credits if credit.role in roles))


def _performing(ctx, credit):
    return credit.role not in ARTIST_ROLES and ctx.tagger.category(credit.role) == PERFORMER_CATEGORY


def _non_performing(ctx, credit):
    return (
        credit.role not in ARTIST_ROLES
        and credit.role not in WRITER_ROLES
        and credit.role != MusicCredit.Role.PRODUCER
        and not _performing(ctx, credit)
    )


def _release_links(track):
    return list(track.release_links.all())


def _releases(track):
    return [link.release for link in _release_links(track)]


def _release_artists(track):
    names = []
    for release in _releases(track):
        names.extend(link.party.display_name for link in release.artist_links.all())
    return ", ".join(_unique(names))


def _companies(track, *kinds):
    return [distribution.company for distribution in track.distributions.all() if distribution.company.kind in kinds]


PRO_NOTE = re.compile(r"PRO:\s*([^\n;]+)", re.IGNORECASE)


def _publisher(company):
    match = PRO_NOTE.search(company.notes or "")
    return f"{company.name} <{match.group(1).strip()}>" if match else company.name


def _links(track, *kinds):
    return [link for link in track.links.all() if not kinds or link.kind in kinds]


def _link_name(link):
    return link.platform or link.name or link.get_kind_display()


def _videos(track):
    return list(track.videos.all())


def _video_urls(track):
    return _join(
        link.url for video in _videos(track) for link in video.links.all() if link.kind == MusicLink.Kind.MUSIC_VIDEO
    )


def _contract_title(contract):
    if contract.titulo:
        return contract.titulo
    asset = contract.file_asset
    return ((asset.attributes or {}).get("name") if asset else "") or "Contract"


def _contract_end(contract):
    return contract.fecha_fin_efectiva or contract.fecha_fin


def _import_value(track, what):
    """File, date or row of the latest import that touched the song."""
    links = list(track.import_links.all())
    if not links:
        return ""
    link = links[0]
    if what == "file":
        return link.import_run.source_name
    if what == "date":
        return _local_date(link.import_run.created_at)
    return link.row_number or ""


def _extra(track):
    extra = (track.metadata or {}).get("extra")
    return extra if isinstance(extra, dict) else {}


def _metadata_id(track_or_release, key):
    ids = (track_or_release.metadata or {}).get("external_ids")
    return ids.get(key, "") if isinstance(ids, dict) else ""


def _credits_and_rights(ctx, track):
    """The original "Credits and rights" cell, with each person's shares."""
    return "; ".join(
        f"{credit.party.display_name} ({ctx.catalog.label(credit.role)}: "
        f"{credit.percentage if credit.percentage is not None else '-'}%, "
        f"publishing {credit.publishing_share if credit.publishing_share is not None else '-'}%)"
        for credit in _credits(track)
    )


def _companies_and_shares(track):
    return "; ".join(
        f"{d.company.name} ({d.company.get_kind_display()}: {d.percentage if d.percentage is not None else '-'}%)"
        for d in track.distributions.all()
    )


# ── Registry ───────────────────────────────────────────────────────────────

C = Column
TrackStatus, TrackKind = MusicTrack.Status.choices, MusicTrack.Kind.choices
ReleaseStatus, ReleaseType = MusicRelease.Status.choices, MusicRelease.Type.choices
Rights = ContractWorkLink.RightsType.choices

COLUMNS = [
    # Canción
    C("title", "Track Name", "song", lambda ctx, t: t.title),
    C("version", "Track Version", "song", lambda ctx, t: t.version),
    C("subtitle", "Subtitle", "song", lambda ctx, t: t.subtitle),
    C("status", "Track Status", "song", lambda ctx, t: _label(TrackStatus, t.status)),
    C("kind", "Track Kind", "song", lambda ctx, t: _label(TrackKind, t.kind)),
    C("release_date", "Track Release Date", "song", lambda ctx, t: t.release_date or ""),
    C("original_release_date", "Original Release Date", "song", lambda ctx, t: t.original_release_date or ""),
    C("recording_date", "Recording Date", "song", lambda ctx, t: t.recording_date or ""),
    C("duration", "Duration", "song", lambda ctx, t: _duration(t.duration_ms)),
    C("genres", "Track Genre", "song", lambda ctx, t: ", ".join(link.genre.name for link in t.genre_links.all()),
      ("genres",)),
    C("language", "Audio Language", "song", lambda ctx, t: t.language),
    C("country_of_recording", "Country Of Recording", "song", lambda ctx, t: t.country_of_recording),
    C("explicit", "Explicit Lyrics", "song", lambda ctx, t: _yes_no(t.explicit)),
    C("instrumental", "Instrumental", "song", lambda ctx, t: _yes_no(t.instrumental)),
    C("digital_format", "Digital Format", "song", lambda ctx, t: t.digital_format),
    C("tiktok_preview", "Sample Start Time", "song", lambda ctx, t: _clock(t.tiktok_preview_start_ms)),
    C("audio_url", "Audio URL", "song", lambda ctx, t: t.audio_url),
    C("cover_url", "Cover URL", "song", lambda ctx, t: t.cover_url),
    # Identificadores
    C("isrc", "Isrc", "identifiers", lambda ctx, t: t.isrc),
    C("isrc_video", "Isrc Video", "identifiers", lambda ctx, t: t.isrc_video),
    C("track_upc", "Track UPC", "identifiers", lambda ctx, t: t.upc),
    C("catalog", "Catalog", "identifiers", lambda ctx, t: t.catalog),
    C("external_id", "Track Id", "identifiers", lambda ctx, t: _metadata_id(t, "track_id")),
    C("release_external_id", "Product Id", "identifiers",
      lambda ctx, t: _one_or_many(_metadata_id(r, "product_id") for r in _releases(t)), ("releases",)),
    C("internal_id", "Internal ID", "identifiers", lambda ctx, t: str(t.id)),
    # Lanzamiento
    C("releases", "Product Title", "release", lambda ctx, t: _join(r.title for r in _releases(t)), ("releases",)),
    C("release_version", "Product Version", "release", lambda ctx, t: _one_or_many(r.version for r in _releases(t)),
      ("releases",)),
    C("release_type", "Release Type", "release",
      lambda ctx, t: _one_or_many(_label(ReleaseType, r.release_type) for r in _releases(t)), ("releases",)),
    C("release_status", "Product Status", "release",
      lambda ctx, t: _one_or_many(_label(ReleaseStatus, r.status) for r in _releases(t)), ("releases",)),
    C("upc", "UPC", "release", lambda ctx, t: _join(_unique(r.upc for r in _releases(t))), ("releases",)),
    C("release_catalog", "Catalog Number", "release",
      lambda ctx, t: _one_or_many(r.catalog_number for r in _releases(t)), ("releases",)),
    C("release_label", "Product Label", "release", lambda ctx, t: _one_or_many(r.label_name for r in _releases(t)),
      ("releases",)),
    C("release_release_date", "Release Date", "release",
      lambda ctx, t: _one_or_many(r.release_date for r in _releases(t)), ("releases",)),
    C("release_original_date", "Product Original Release Date", "release",
      lambda ctx, t: _one_or_many(r.original_release_date for r in _releases(t)), ("releases",)),
    C("release_language", "Primary Metadata Language", "release",
      lambda ctx, t: _one_or_many(r.language for r in _releases(t)), ("releases",)),
    C("release_artists", "Product Display Artist", "release", lambda ctx, t: _release_artists(t), ("releases",)),
    C("disc_number", "Track Volume", "release",
      lambda ctx, t: _one_or_many(link.disc_number for link in _release_links(t)), ("releases",)),
    C("track_number", "Track Number", "release",
      lambda ctx, t: _one_or_many(link.track_number for link in _release_links(t)), ("releases",)),
    C("release_p_line", "Product P Line", "release", lambda ctx, t: _one_or_many(r.p_line for r in _releases(t)),
      ("releases",)),
    C("release_c_line", "Product C Line", "release", lambda ctx, t: _one_or_many(r.c_line for r in _releases(t)),
      ("releases",)),
    # Artistas y créditos
    C("artists", "Track Display Artist", "credits",
      lambda ctx, t: _names(_credits(t), (MusicCredit.Role.PRIMARY_ARTIST,)), ("credits",)),
    C("featured_artists", "Featured Artists", "credits",
      lambda ctx, t: _names(_credits(t), (MusicCredit.Role.FEATURED_ARTIST,)), ("credits",)),
    C("performing_contributors", "Performing Contributors", "credits",
      lambda ctx, t: _tagged(ctx, [c for c in _credits(t) if _performing(ctx, c)]), ("credits",)),
    C("non_performing_contributors", "Non-Performing Contributors", "credits",
      lambda ctx, t: _tagged(ctx, [c for c in _credits(t) if _non_performing(ctx, c)]), ("credits",)),
    C("writers", "Writers", "credits",
      lambda ctx, t: _tagged(ctx, [c for c in _credits(t) if c.role in WRITER_ROLES]), ("credits",)),
    C("producers", "Producers", "credits", lambda ctx, t: _names(_credits(t), (MusicCredit.Role.PRODUCER,)),
      ("credits",)),
    C("contributors_all", "Contributors (all roles)", "credits", lambda ctx, t: _tagged(ctx, _credits(t)),
      ("credits",)),
    C("credits_rights", "Credits and rights", "credits", lambda ctx, t: _credits_and_rights(ctx, t), ("credits",)),
    # Compañías y porcentajes
    C("publishers", "Publishers/Collection Societies", "companies",
      lambda ctx, t: ", ".join(_publisher(c) for c in _companies(t, MusicCompany.Kind.PUBLISHER)),
      ("distributions",)),
    C("record_labels", "Label", "companies",
      lambda ctx, t: ", ".join(c.name for c in _companies(t, MusicCompany.Kind.RECORD_LABEL)), ("distributions",)),
    C("aggregators", "Aggregator", "companies",
      lambda ctx, t: ", ".join(c.name for c in _companies(t, MusicCompany.Kind.AGGREGATOR)), ("distributions",)),
    C("distributors", "Distributor", "companies",
      lambda ctx, t: ", ".join(c.name for c in _companies(t, MusicCompany.Kind.DISTRIBUTOR)), ("distributions",)),
    C("other_companies", "Other Companies", "companies",
      lambda ctx, t: ", ".join(
          f"{c.name} <{c.get_kind_display()}>"
          for c in _companies(t, MusicCompany.Kind.MANAGEMENT, MusicCompany.Kind.OTHER)
      ),
      ("distributions",)),
    C("companies_shares", "Companies and shares", "companies", lambda ctx, t: _companies_and_shares(t),
      ("distributions",)),
    C("aggregator_percentage", "Aggregator Percentage", "companies",
      lambda ctx, t: _number(t.aggregator_percentage)),
    C("distributor_percentage", "Distributor Percentage", "companies",
      lambda ctx, t: _number(t.distributor_percentage)),
    C("record_label_percentage", "Record Label Percentage", "companies",
      lambda ctx, t: _number(t.record_label_percentage)),
    C("artist_percentage", "Artist Percentage", "companies", lambda ctx, t: _number(t.artist_percentage)),
    C("writer_percentage", "Writer Percentage", "companies", lambda ctx, t: _number(t.writer_percentage)),
    # Videos y enlaces
    C("videos", "Videos", "media", lambda ctx, t: _aligned(v.title for v in _videos(t)), ("videos",)),
    C("video_isrcs", "Video ISRCs", "media", lambda ctx, t: _aligned(v.isrc_video or v.isrc for v in _videos(t)),
      ("videos",)),
    C("video_upcs", "Video UPCs", "media", lambda ctx, t: _aligned(v.upc for v in _videos(t)), ("videos",)),
    C("video_dates", "Video Release Date", "media", lambda ctx, t: _aligned(v.release_date for v in _videos(t)),
      ("videos",)),
    C("video_urls", "Video URL", "media", lambda ctx, t: _video_urls(t), ("videos",)),
    C("streaming_links", "Track Play Link", "media",
      lambda ctx, t: _join(link.url for link in _links(t, MusicLink.Kind.STREAMING)), ("links",)),
    C("store_links", "Store Links", "media",
      lambda ctx, t: _join(link.url for link in _links(t, MusicLink.Kind.STORE)), ("links",)),
    C("social_links", "Social Links", "media",
      lambda ctx, t: _join(link.url for link in _links(t, MusicLink.Kind.SOCIAL)), ("links",)),
    C("other_links", "Other Links", "media",
      lambda ctx, t: _join(link.url for link in _links(t, MusicLink.Kind.OTHER)), ("links",)),
    C("links_all", "Links (all)", "media",
      lambda ctx, t: _join(f"{_link_name(link)}: {link.url}" for link in _links(t)),
      ("links",)),
    # Derechos y textos
    C("p_line", "Track P Line", "rights", lambda ctx, t: t.p_line),
    C("c_line", "Track C Line", "rights", lambda ctx, t: t.c_line),
    C("ownership", "Ownership", "rights", lambda ctx, t: t.ownership),
    C("us_publishing_obligations", "US Publishing Obligations", "rights", lambda ctx, t: t.us_publishing_obligations),
    C("recoupment", "Recoupment", "rights", lambda ctx, t: t.recoupment),
    C("lyrics", "Lyrics", "rights", lambda ctx, t: t.lyrics),
    # Contratos (linked to the song, or to one of its releases)
    C("contracts", "Contracts", "contracts",
      lambda ctx, t: _aligned(_contract_title(c) for c, _r in ctx.contracts.get(t.id, ())), ("contracts",)),
    C("contract_status", "Contract Status", "contracts",
      lambda ctx, t: _aligned(c.get_estatus_contrato_display() if c.estatus_contrato else ""
                              for c, _r in ctx.contracts.get(t.id, ())),
      ("contracts",)),
    C("contract_end", "Contract End Date", "contracts",
      lambda ctx, t: _aligned(_contract_end(c) for c, _r in ctx.contracts.get(t.id, ())), ("contracts",)),
    C("contract_rights", "Contract Rights", "contracts",
      lambda ctx, t: _aligned(_label(Rights, r) for _c, r in ctx.contracts.get(t.id, ())), ("contracts",)),
    # Importación y datos extra
    C("import_file", "Import File", "import", lambda ctx, t: _import_value(t, "file"), ("imports",)),
    C("import_date", "Import Date", "import", lambda ctx, t: _import_value(t, "date"), ("imports",)),
    C("import_row", "Import Row", "import", lambda ctx, t: _import_value(t, "row"), ("imports",)),
    C(EXTRA, EXTRA_PREFIX + "…", "import", lambda ctx, t: ""),
    C("created_at", "Created At", "import", lambda ctx, t: _local_date(t.created_at)),
    C("updated_at", "Updated At", "import", lambda ctx, t: _local_date(t.updated_at)),
    # Etiquetas
    C("tags", "Tags", "tags", lambda ctx, t: ", ".join(ctx.tags.get(t.id, ())), ("tags",)),
]
REGISTRY = {column.key: column for column in COLUMNS}
assert len(REGISTRY) == len(COLUMNS), "duplicate export column key"
assert {column.group for column in COLUMNS} <= set(GROUPS)

# The original export (no `columns`): these columns with their old headers.
LEGACY_HEADERS = {
    "title": "Song",
    "version": "Version",
    "artists": "Artists",
    "featured_artists": "Featured artists",
    "isrc": "ISRC",
    "status": "Status",
    "release_date": "Release date",
    "releases": "Releases",
    "upc": "UPC",
    "genres": "Genres",
    "videos": "Videos",
    "video_isrcs": "Video ISRCs",
    "video_dates": "Video dates",
    "video_urls": "Video URLs",
    "credits_rights": "Credits and rights",
    "companies_shares": "Companies and shares",
    "ownership": "Ownership",
    "p_line": "P line",
    "c_line": "C line",
}
BASIC = list(LEGACY_HEADERS)


def parse_columns(raw):
    """(keys, legacy) for a ``columns`` parameter: missing -> the original
    columns with their old headers, "all" -> every column, otherwise the
    known keys in the order given. Raises ValueError when none is known."""
    text = str(raw or "").strip()
    if not text:
        return BASIC, True
    if text.lower() == "all":
        return [column.key for column in COLUMNS], False
    keys = []
    for key in text.split(","):
        key = key.strip()
        if key in REGISTRY and key not in keys:
            keys.append(key)
    if not keys:
        raise ValueError("No known export columns")
    return keys, False


# ── Building the file ──────────────────────────────────────────────────────


def _prefetch(queryset, needs):
    lookups = []
    if "credits" in needs:
        lookups.append(
            Prefetch(
                "credits",
                queryset=MusicCredit.objects.select_related("party").order_by("created_at", "id"),
            )
        )
    if "releases" in needs:
        lookups.append(
            Prefetch(
                "release_links",
                queryset=MusicReleaseTrack.objects.select_related("release")
                .prefetch_related(
                    Prefetch(
                        "release__artist_links",
                        queryset=MusicReleaseArtist.objects.select_related("party").order_by("role", "created_at"),
                    )
                )
                .order_by("release__release_date", "disc_number", "track_number"),
            )
        )
    if "genres" in needs:
        lookups.append(Prefetch("genre_links", queryset=MusicTrackGenre.objects.select_related("genre")))
    if "distributions" in needs:
        lookups.append(
            Prefetch(
                "distributions", queryset=MusicDistribution.objects.select_related("company").order_by("created_at")
            )
        )
    if "videos" in needs:
        lookups.append(
            Prefetch(
                "videos",
                queryset=MusicTrack.objects.order_by("release_date", "created_at").prefetch_related(
                    Prefetch("links", queryset=MusicLink.objects.order_by("created_at"))
                ),
            )
        )
    if "links" in needs:
        lookups.append(Prefetch("links", queryset=MusicLink.objects.order_by("kind", "created_at")))
    if "imports" in needs:
        lookups.append(
            Prefetch(
                "import_links",
                queryset=MusicTrackImport.objects.select_related("import_run").order_by("-created_at"),
            )
        )
    return queryset.prefetch_related(*lookups) if lookups else queryset


def _contracts_by_track(workspace_id, tracks):
    """Track id -> [(contract, rights_type)] for LINKED contracts of the song
    or of any of its releases, each contract once (direct links first)."""
    track_ids = [track.id for track in tracks]
    if not track_ids:
        return {}
    releases_by_track = {}
    for track_id, release_id in MusicReleaseTrack.objects.filter(track_id__in=track_ids).values_list(
        "track_id", "release_id"
    ):
        releases_by_track.setdefault(track_id, []).append(release_id)
    release_ids = {release_id for ids in releases_by_track.values() for release_id in ids}
    links = (
        ContractWorkLink.objects.filter(
            workspace_id=workspace_id,
            status=ContractWorkLink.Status.LINKED,
            contract__file_asset__is_deleted=False,
        )
        .filter(
            Q(entity_type=ContractWorkLink.EntityType.MUSIC_TRACK, entity_id__in=track_ids)
            | Q(entity_type=ContractWorkLink.EntityType.MUSIC_RELEASE, entity_id__in=release_ids)
        )
        .select_related("contract__file_asset")
        .order_by("created_at")
    )
    direct, by_release = {}, {}
    for link in links:
        target = direct if link.entity_type == ContractWorkLink.EntityType.MUSIC_TRACK else by_release
        target.setdefault(link.entity_id, []).append(link)
    result = {}
    for track_id in track_ids:
        entries, seen = [], set()
        candidates = list(direct.get(track_id, ()))
        for release_id in releases_by_track.get(track_id, ()):
            candidates.extend(by_release.get(release_id, ()))
        for link in candidates:
            if link.contract_id in seen:
                continue
            seen.add(link.contract_id)
            entries.append((link.contract, link.rights_type))
        if entries:
            result[track_id] = entries
    return result


def _tags_by_track(workspace_id, tracks):
    result = {}
    rows = (
        EntityTagLink.objects.filter(
            workspace_id=workspace_id,
            entity_type="music_track",
            entity_id__in=[track.id for track in tracks],
            status=EntityTagLink.Status.CONFIRMED,
            tag__deleted_at__isnull=True,
        )
        .order_by("tag__name")
        .values_list("entity_id", "tag__name")
    )
    for entity_id, name in rows:
        result.setdefault(entity_id, []).append(name)
    return result


def workspace_extra_keys(workspace_id, limit=500):
    """Every key songs of the workspace keep in metadata["extra"]."""
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT key FROM (
                SELECT DISTINCT jsonb_object_keys(metadata->'extra') AS key
                FROM music_tracks
                WHERE workspace_id = %s AND deleted_at IS NULL AND parent_track_id IS NULL
                  AND jsonb_typeof(metadata->'extra') = 'object'
            ) keys ORDER BY lower(key) LIMIT %s
            """,
            [str(workspace_id), limit],
        )
        return [row[0] for row in cursor.fetchall()]


def describe_columns(workspace):
    """The registry for the export panel."""
    return {
        "groups": list(GROUPS),
        "columns": [{"key": column.key, "header": column.header, "group": column.group} for column in COLUMNS],
        "presets": {"basic": BASIC, "all": [column.key for column in COLUMNS]},
        "extra_keys": workspace_extra_keys(workspace.id),
    }


def build_rows(workspace, queryset, keys, legacy=False):
    """(headers, rows) for the songs in `queryset` and the chosen columns."""
    columns = [REGISTRY[key] for key in keys]
    needs = {need for column in columns for need in column.needs}
    tracks = list(_prefetch(queryset, needs))
    ctx = ExportContext(
        tagger=workspace_credit_tagger(workspace.id) if "credits" in needs else CreditTagger(),
        catalog=workspace_role_catalog(workspace.id) if "credits_rights" in keys else RoleCatalog(),
    )
    if "contracts" in needs:
        ctx.contracts = _contracts_by_track(workspace.id, tracks)
    if "tags" in needs:
        ctx.tags = _tags_by_track(workspace.id, tracks)
    if EXTRA in keys:
        seen = {}
        for track in tracks:
            for key in _extra(track):
                seen.setdefault(str(key), None)
        ctx.extra_keys = list(seen)

    headers = []
    for column in columns:
        if column.key == EXTRA:
            headers.extend(f"{EXTRA_PREFIX}{key}" for key in ctx.extra_keys)
        else:
            headers.append(LEGACY_HEADERS.get(column.key, column.header) if legacy else column.header)
    rows = []
    for track in tracks:
        row = []
        for column in columns:
            if column.key == EXTRA:
                extra = _extra(track)
                row.extend(_cell(extra.get(key, "")) for key in ctx.extra_keys)
            else:
                row.append(_cell(column.value(ctx, track)))
        rows.append(row)
    return headers, rows


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return str(value)
    return value


def _excel_value(value):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > EXCEL_CELL_LIMIT:
            value = value[: EXCEL_CELL_LIMIT - 1] + "…"
    return value


def _excel(headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Music catalog"
    sheet.append(headers)
    bold = Font(bold=True)
    for cell in sheet[1]:
        cell.font = bold
    widths = [len(str(header)) for header in headers]
    for index, row in enumerate(rows):
        values = [_excel_value(value) for value in row]
        sheet.append(values)
        line = index + 2
        for position, value in enumerate(values):
            if isinstance(value, str):
                # "=..." stays text (a UPC like ="0198…"), never a formula
                if value.startswith("="):
                    cell = sheet.cell(row=line, column=position + 1)
                    cell.data_type = "s"
                if index < 500:
                    widths[position] = max(widths[position], min(len(value), 60))
            elif index < 500 and value not in (None, ""):
                widths[position] = max(widths[position], len(_text(value)))
    for position, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = max(8, min(width + 2, 60))
    sheet.freeze_panes = "A2"
    if headers:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def report_response(headers, rows, export_format, filename):
    if export_format == "xlsx":
        response = HttpResponse(
            _excel(headers, rows),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return response
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    response.write("﻿")
    writer = csv.writer(response)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_text(value) if isinstance(value, (date, datetime)) else value for value in row])
    return response
